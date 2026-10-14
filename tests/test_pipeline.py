"""CPU-only regression tests. No model downloads or benchmark labels in selection."""
import gzip
import json
import tempfile
import unittest
import sys
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from spreadsheet_agent.consistency import select, strict_fingerprint
from spreadsheet_agent.io import ROOT, atomic, file_hash
from spreadsheet_agent.predict import routes, select_outputs
from spreadsheet_agent import predict
from spreadsheet_agent.review import decode, correction_prompt
from spreadsheet_agent.reward_terms import score_record
from spreadsheet_agent import domain_preflight, sandbox
from spreadsheet_code_agent_v0.preflight import run_preflight
from spreadsheet_code_agent_v0 import run_agent_v0 as agent
from spreadsheetbench_lite.data import load_tasks
from spreadsheetbench_lite.llm import LLMResult
from spreadsheet_code_agent_v0.runtime import AgentExecResult, verify_output
from spreadsheet_recalc.evaluate import inventory


class Tokenizer:
    def encode(self, text, **kwargs):
        return list(range(len(text) // 4))


def workbook(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active["A1"] = value
    wb.save(path)
    wb.close()


class PipelineTests(unittest.TestCase):
    def test_frozen_recipes(self):
        meta = json.loads((ROOT / "provenance/training.json").read_text())
        for stage, expected in meta.items():
            recipe = ROOT / "recipes" / f"{stage}.jsonl.gz"
            self.assertEqual(file_hash(recipe), expected["recipe_sha256"])
            rows = [json.loads(line) for line in gzip.decompress(recipe.read_bytes()).decode().splitlines()]
            self.assertEqual(len(rows), expected["rows"])
            self.assertEqual(len({r["extra_info"]["task_id"] for r in rows}), expected["unique_tasks"])
            self.assertTrue(all("dataset_dir" not in r["extra_info"] for r in rows))
            if stage == "review":
                self.assertTrue(all(r["extra_info"]["baseline_code"] for r in rows))

    def test_routes(self):
        self.assertEqual(routes("mixed"), [("common",) * 3] * 2 + [("base", "generation", "common")] * 2)
        self.assertEqual(routes("free"), [("base",) * 3] * 4)

    def test_selection_ties_and_missing(self):
        obs = lambda fp, ok=True: dict(executable=ok, fingerprint=fp)
        self.assertEqual(select([obs("a"), obs("b"), obs("b"), obs("a")])["consensus"], 0)
        self.assertEqual(select([obs(None, False), obs("b"), obs("c")])["consensus"], 1)
        self.assertEqual(select([obs(None, False)] * 4)["consensus"], 0)
        self.assertEqual(select([obs("a"), obs("b"), obs("b")])["consensus"], 1)

    def test_formula_not_empty_cache(self):
        with tempfile.TemporaryDirectory() as d:
            a, b = Path(d) / "a.xlsx", Path(d) / "b.xlsx"
            workbook(a, "=1+1")
            workbook(b, "=2+2")
            self.assertNotEqual(strict_fingerprint(a), strict_fingerprint(b))

    def test_domain_fix_is_separate(self):
        code = 'def transform(wb):\n    wb.create_sheet("New")\n    wb["New"]["A1"] = 2\n'
        inspection = dict(sheet_names=["Sheet"], sheets=[])
        original = run_preflight(code, inspection)
        self.assertTrue(any(e.category == "sheet_error" for e in original.errors))
        fixed = domain_preflight.run_preflight(code, inspection)
        self.assertFalse(any(e.category == "sheet_error" for e in fixed.errors))
        self.assertEqual(run_preflight(code, inspection).to_dict(), original.to_dict())

    def test_review_actions(self):
        original = 'def transform(wb):\n    wb.active["A1"] = 2'
        self.assertEqual(decode("DECISION: KEEP", original), ("KEEP", original))
        action, code = decode("DECISION: REVISE\n```python\n" + original + "\n```", original)
        self.assertEqual(action, "REVISE")
        self.assertIn("transform(wb)", code)
        with self.assertRaises(ValueError):
            decode("Looks good", original)
        self.assertIn("VISIBLE INPUT/OUTPUT", correction_prompt("Add one", original, {}))

    def test_rewards(self):
        wrong = {"cases": [{"ok": True, "verification": {}}], "eval": {"hard_restriction": 0, "soft_restriction": 0}}
        right = {"cases": [{"ok": True, "verification": {}}], "eval": {"hard_restriction": 1, "soft_restriction": 1}}
        for strategy in ("diagnostic", "diagnostic_aligned"):
            self.assertGreater(score_record(right, "", strategy)["scalar"], score_record(wrong, "", strategy)["scalar"])
        self.assertEqual(score_record(right, "", "diagnostic_aligned", {"score": 1, "exact": True})["scalar"], 1)

    def test_selection_and_recalc_inventory_both_datasets(self):
        # Candidate 0 is labelled correct, but candidates 1/2 agree. Selection
        # must choose 1 even though reading the labels could yield a better score.
        for prefix, variants in (("v1", 3), ("domain", 1)):
            with self.subTest(dataset=prefix), tempfile.TemporaryDirectory() as d:
                root = Path(d)
                dataset, out = root / "dataset", root / "selected"
                ident = f"{prefix}_toy"
                atomic(dataset / "dataset.json", [dict(id=ident, instruction="Write two to A1", answer_position="Sheet!A1",
                       spreadsheet_path=f"spreadsheet/{ident}")])
                for case in range(1, variants + 1):
                    for kind, value in (("input", 0), ("answer", 2)):
                        workbook(dataset / "spreadsheet" / ident / f"{case}_{ident}_{kind}.xlsx", value)
                roots = []
                for index, value in enumerate([2, 3, 3, 4]):
                    folder = root / f"candidate{index}"
                    roots.append(folder)
                    cases = []
                    for case in range(1, variants + 1):
                        target = folder / "spreadsheet" / ident / f"{case}_{ident}_output.xlsx"
                        workbook(target, value)
                        cases.append(dict(case=case, ok=True, output=str(target)))
                    (folder / "codes").mkdir()
                    (folder / "codes" / f"{ident}.py").write_text(f"def transform(wb):\n    wb.active['A1'] = {value}\n")
                    record = dict(id=ident, cases=cases, ok=True, eval=dict(test_case_results=[int(index == 0)] * variants,
                                  hard_restriction=int(index == 0), soft_restriction=float(index == 0)))
                    (folder / "run_results.jsonl").write_text(json.dumps(record) + "\n")
                select_outputs(load_tasks(dataset), dataset, roots, out, Tokenizer())
                self.assertEqual(json.loads((out / "choices" / f"{ident}.json").read_text())["selected"], 1)
                self.assertEqual(json.loads((out / "summary.json").read_text())["metrics"]["hard_accuracy"], 0)
                self.assertEqual(len(inventory(dataset, out, 0, None)), variants)

    def test_sandbox_fails_closed(self):
        with patch("spreadsheet_agent.sandbox.shutil.which", return_value=None):
            with self.assertRaises(sandbox.InfrastructureError):
                sandbox.sandbox_command(Path("/tmp"))

    def test_final_preflight_removes_stale_output(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "output.xlsx"
            workbook(path, 2)
            task = SimpleNamespace(id="toy", instruction_type="", answer_position="Sheet!A1")
            case = SimpleNamespace(index=1, input_path=Path(d) / "input.xlsx", output_path=path)
            args = SimpleNamespace(max_repairs=0)
            record, code = agent._run_with_optional_repair(task, [case], "def transform(wb): !!!", args,
                {"repair": False, "preflight": True}, {}, {}, "", "", None, {}, None, {})
            self.assertFalse(record["ok"])
            self.assertFalse(path.exists())

    def test_mocked_full_prediction_and_resume(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            dataset = root / "dataset"
            atomic(dataset / "dataset.json", [dict(id="toy", instruction="Write 2 to A1", answer_position="Sheet!A1",
                   spreadsheet_path="spreadsheet/toy")])
            for kind, value in (("input", 0), ("answer", 2)):
                workbook(dataset / "spreadsheet/toy" / f"1_toy_{kind}.xlsx", value)
            calls = []
            def completion(prompt, config):
                calls.append(config.model)
                content = json.dumps(dict(where="Sheet!A1", what="Set A1 to 2", **{"from": ""}, risks=[], api_needs=[])) if config.temperature == 0 else '```python\ndef transform(wb):\n    wb["Sheet"]["A1"] = 2\n```'
                return LLMResult(content, "", "stop", {}, {})
            def trusted_execution(code, source, destination, timeout, position):
                # Execute only the fixed test fixture, never arbitrary model code.
                self.assertIn('wb["Sheet"]["A1"] = 2', code)
                wb = load_workbook(source)
                wb.active["A1"] = 2
                destination.parent.mkdir(parents=True, exist_ok=True)
                wb.save(destination)
                wb.close()
                return AgentExecResult(True, 0, "", "", False, "", verify_output(source, destination, position))
            models = SimpleNamespace(list=lambda: SimpleNamespace(model_dump=lambda **k: {
                "data": [dict(id=alias) for alias in ("base", "common", "generation")]}))
            client = lambda **k: SimpleNamespace(models=models)
            argv = ["predict", "--dataset", "v1", "--dataset-dir", str(dataset), "--out", str(root / "run"),
                    "--tokenizer", "fixture", "--limit", "1"]
            with patch("openai.OpenAI", client), patch("transformers.AutoTokenizer.from_pretrained", return_value=Tokenizer()), \
                 patch.object(agent, "chat_completion", completion), patch.object(sandbox, "execute", trusted_execution), \
                 patch.object(agent, "run_transform_code"), patch.object(sys, "argv", argv):
                predict.main()
                self.assertEqual(len(calls), 8)
                self.assertEqual(calls, ["common", "common", "common", "common", "base", "generation", "base", "generation"])
                # The legacy agent rewrites sys.argv internally; restore the entry args.
                sys.argv = argv
                predict.main()
                self.assertEqual(len(calls), 8)
            summary = json.loads((root / "run/selected/summary.json").read_text())
            self.assertEqual(summary["metrics"]["hard_accuracy"], 1)
            self.assertEqual(summary["metrics"]["average_accuracy"], 1)


if __name__ == "__main__":
    unittest.main()
