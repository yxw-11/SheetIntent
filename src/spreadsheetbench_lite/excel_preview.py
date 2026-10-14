from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def workbook_preview(path: Path, rows: int = 5, max_cols: int = 30) -> str:
    """Return a compact, deterministic text preview of every sheet.

    SpreadsheetBench V1 official single-round prompts provide the first few rows.
    This version uses openpyxl instead of pandas so the runner stays small.
    """

    workbook = load_workbook(path, data_only=False)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        chunks.append(f"Sheet Name: {sheet.title}")
        chunks.append(f"Dimensions: {sheet.max_row} rows x {sheet.max_column} columns")
        if sheet.merged_cells.ranges:
            merged = ", ".join(str(cell_range) for cell_range in sheet.merged_cells.ranges)
            chunks.append(f"Merged Cells: {merged}")
        chunks.extend(_sheet_rows(sheet.iter_rows(max_row=rows, max_col=min(sheet.max_column, max_cols))))
        chunks.append("-" * 50)
    return "\n".join(chunks)


def _sheet_rows(rows: Iterable[tuple]) -> list[str]:
    rendered = ["Preview:"]
    for row in rows:
        values = [_format_cell(cell.value) for cell in row]
        rendered.append("\t".join(values).rstrip())
    return rendered


def _format_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\t", " ")
    if len(text) > 120:
        text = text[:117] + "..."
    return text
