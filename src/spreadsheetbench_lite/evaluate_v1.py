from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils.cell import get_column_letter, range_boundaries
from tqdm import tqdm

from .data import discover_test_cases, load_tasks


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Lightweight SpreadsheetBench V1 answer-position evaluator.")
    parser.add_argument("--dataset-dir", type=Path, required=True)
    parser.add_argument("--outputs-dir", type=Path, required=True, help="Directory containing <task_id>/<case>_<id>_output.xlsx.")
    parser.add_argument("--result-file", type=Path, default=None)
    parser.add_argument("--start", type=int, default=0, help="Start offset in dataset.json.")
    parser.add_argument("--limit", type=int, default=None, help="Optional number of tasks to evaluate.")
    parser.add_argument("--run-results", type=Path, default=None, help="Evaluate only ids present in a run_results.jsonl file.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    dataset_dir = args.dataset_dir.resolve()
    outputs_dir = args.outputs_dir.resolve()
    result_file = args.result_file or outputs_dir.parent / "eval_results.json"

    tasks = load_tasks(dataset_dir)
    if args.run_results:
        wanted_ids = _ids_from_run_results(args.run_results)
        tasks = [task for task in tasks if task.id in wanted_ids]
    else:
        tasks = tasks[args.start :]
        if args.limit is not None:
            tasks = tasks[: args.limit]

    records = []
    for task in tqdm(tasks, desc="evaluate"):
        cases = discover_test_cases(dataset_dir, outputs_dir.parent, task)
        case_scores = []
        for case in cases:
            output_path = outputs_dir / task.id / case.output_path.name
            ok, message = compare_workbooks(case.answer_path, output_path, task.answer_position)
            case_scores.append({"case": case.index, "ok": int(ok), "message": message})
        hard = int(bool(case_scores) and all(item["ok"] for item in case_scores))
        soft = sum(item["ok"] for item in case_scores) / len(case_scores) if case_scores else 0.0
        records.append(
            {
                "id": task.id,
                "instruction_type": task.instruction_type,
                "test_case_results": [item["ok"] for item in case_scores],
                "soft_restriction": soft,
                "hard_restriction": hard,
                "messages": case_scores,
            }
        )

    result_file.parent.mkdir(parents=True, exist_ok=True)
    result_file.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    total = len(records)
    hard_acc = sum(item["hard_restriction"] for item in records) / total if total else 0.0
    soft_acc = sum(item["soft_restriction"] for item in records) / total if total else 0.0
    print(json.dumps({"tasks": total, "hard_accuracy": hard_acc, "soft_accuracy": soft_acc}, indent=2))
    print(f"wrote {result_file}")


def _ids_from_run_results(path: Path) -> set[str]:
    ids = set()
    with path.open("r", encoding="utf-8") as fp:
        for line in fp:
            if line.strip():
                ids.add(str(json.loads(line)["id"]))
    return ids


def compare_workbooks(answer_path: Path | None, output_path: Path, answer_position: str) -> tuple[bool, str]:
    if answer_path is None or not answer_path.exists():
        return False, "answer workbook not found"
    if not output_path.exists():
        return False, "output workbook not found"

    try:
        answer_wb = openpyxl.load_workbook(answer_path, data_only=True)
        output_wb = openpyxl.load_workbook(output_path, data_only=True)
    except Exception as exc:
        return False, str(exc)

    if answer_position in ("", "__FULL_WORKBOOK__"):
        return compare_full_workbook(answer_wb, output_wb)

    try:
        positions = parse_answer_positions(answer_wb, answer_position)
    except Exception as exc:
        return False, f"invalid answer_position {answer_position!r}: {exc}"

    for sheet_name, cell_range in positions:
        if sheet_name not in answer_wb.sheetnames:
            return False, f"answer worksheet not found: {sheet_name}"
        if sheet_name not in output_wb.sheetnames:
            return False, f"worksheet not found: {sheet_name}"
        answer_ws = answer_wb[sheet_name]
        output_ws = output_wb[sheet_name]
        try:
            cell_names = expand_cell_range(cell_range, answer_ws)
        except Exception as exc:
            return False, f"invalid cell range {cell_range!r} from {answer_position!r}: {exc}"
        for cell_name in cell_names:
            if not compare_value(answer_ws[cell_name].value, output_ws[cell_name].value):
                return False, (
                    f"{sheet_name}!{cell_name}: answer={answer_ws[cell_name].value!r}, "
                    f"output={output_ws[cell_name].value!r}"
                )
    return True, ""


def compare_full_workbook(answer_wb: openpyxl.Workbook, output_wb: openpyxl.Workbook) -> tuple[bool, str]:
    """Internal helper for RL training tasks that do not expose answer_position.

    SpreadsheetBench V1 public evaluation still uses explicit answer ranges; this
    full-workbook mode is intended for converted Spreadsheet-RL training data.
    """

    for sheet_name in answer_wb.sheetnames:
        if sheet_name not in output_wb.sheetnames:
            return False, f"worksheet not found: {sheet_name}"
        answer_ws = answer_wb[sheet_name]
        output_ws = output_wb[sheet_name]
        for row in range(1, answer_ws.max_row + 1):
            for col in range(1, answer_ws.max_column + 1):
                cell_name = f"{get_column_letter(col)}{row}"
                if not compare_value(answer_ws.cell(row, col).value, output_ws.cell(row, col).value):
                    return False, (
                        f"{sheet_name}!{cell_name}: answer={answer_ws.cell(row, col).value!r}, "
                        f"output={output_ws.cell(row, col).value!r}"
                    )
    return True, ""


def parse_answer_positions(workbook: openpyxl.Workbook, answer_position: str) -> list[tuple[str, str]]:
    default_sheet = workbook.sheetnames[0]
    positions = []
    for item in _answer_position_items(answer_position):
        item = item.strip()
        if not item:
            continue
        sheet_name, cell_range = _split_answer_position_item(item, default_sheet)
        positions.append(
            (
                _normalize_sheet_name(sheet_name, workbook.sheetnames),
                _normalize_cell_range_text(cell_range),
            )
        )
    return positions


def _answer_position_items(value: str) -> list[str]:
    # Excel's tokenizer preserves commas and escaped quotes inside sheet names.
    try:
        tokens = Tokenizer("=" + value).items
    except TokenizerError:
        return value.split(",")  # Retain support for legacy malformed quoting.
    items = [""]
    for token in tokens:
        if token.value == "," and token.type == "OPERATOR-INFIX":
            items.append("")
        else:
            items[-1] += token.value
    return items


def _split_answer_position_item(item: str, default_sheet: str) -> tuple[str, str]:
    # SpreadsheetBench metadata contains a few malformed-but-readable forms,
    # e.g. Sheet1'!A1:B2 and 'Calculation!'C11:I490'. Normalize those here.
    item = item.strip()
    match = re.match(r"^'([^']+)!'(.+)'$", item)
    if match:
        return match.group(1), match.group(2)
    match = re.match(r"^'((?:[^']|'')+)'!(.+)$", item)
    if match:
        return match.group(1), match.group(2)
    if "!" in item:
        sheet_name, cell_range = item.split("!", 1)
        return _strip_excel_quotes(sheet_name), _strip_excel_quotes(cell_range)
    return default_sheet, item


def _normalize_sheet_name(value: str, sheetnames: list[str]) -> str:
    if value in sheetnames:
        return value
    value = _strip_excel_quotes(value)
    candidates = [
        value,
        value.strip("'\" "),
        value.rstrip("'\" ").lstrip("'\" "),
        value.replace("''", "'"),
    ]
    for candidate in candidates:
        if candidate in sheetnames:
            return candidate
    lowered = {name.lower(): name for name in sheetnames}
    for candidate in candidates:
        found = lowered.get(candidate.lower())
        if found:
            return found
    matches = [name for name in sheetnames
               if any(name.strip().casefold() == candidate.strip().casefold() for candidate in candidates)]
    if not matches:
        matches = [name for name in sheetnames
                   if any(" ".join(name.split()).casefold() == " ".join(candidate.split()).casefold()
                          for candidate in candidates)]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError(f"ambiguous worksheet name: {value!r}")
    return candidates[1] if len(candidates) > 1 else value


def expand_cell_range(cell_range: str, worksheet: Any | None = None) -> list[str]:
    cell_range = _normalize_cell_range_text(_drop_sheet_prefix(cell_range))
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    if min_col is None:
        min_col = 1
    if max_col is None:
        max_col = worksheet.max_column if worksheet is not None else min_col
    if min_row is None:
        min_row = 1
    if max_row is None:
        max_row = worksheet.max_row if worksheet is not None else min_row
    start_col, end_col = sorted((min_col, max_col))
    start_row, end_row = sorted((min_row, max_row))
    return [
        f"{get_column_letter(col)}{row}"
        for col in range(start_col, end_col + 1)
        for row in range(start_row, end_row + 1)
    ]


def _strip_excel_quotes(value: str) -> str:
    value = value.strip()
    value = value.strip("'\" ")
    while len(value) >= 2 and value[0] == "'" and value[-1] == "'":
        value = value[1:-1].strip()
    return value


def _normalize_cell_range_text(value: str) -> str:
    value = _strip_excel_quotes(value).replace("：", ":").replace(" ", "")
    value = value.strip("'\" ")
    match = re.match(r"^([A-Za-z]+)(\d+):(\d+)$", value)
    if match:
        col, start_row, end_row = match.groups()
        return f"{col}{start_row}:{col}{end_row}"
    return value


def _drop_sheet_prefix(cell_range: str) -> str:
    cell_range = _strip_excel_quotes(cell_range)
    if "!" in cell_range:
        cell_range = cell_range.rsplit("!", 1)[1]
    return _strip_excel_quotes(cell_range)


def compare_value(left: Any, right: Any) -> bool:
    left = normalize_value(left)
    right = normalize_value(right)
    if left in ("", None) and right in ("", None):
        return True
    return type(left) is type(right) and left == right


def normalize_value(value: Any) -> Any:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, dt.datetime):
        base = dt.datetime(1899, 12, 30)
        delta = value - base
        return round(delta.days + delta.seconds / 86400.0, 0)
    if isinstance(value, dt.time):
        return str(value)[:-3]
    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


if __name__ == "__main__":
    main()
