from __future__ import annotations

import hashlib
import json
import os
import re
import time
from pathlib import Path
from urllib.parse import urlsplit

import httpx
import openpyxl


class RecalculationError(RuntimeError):
    """Infrastructure failure, never a model correctness label."""


def digest(path: Path) -> str | None:
    if not path.exists():
        return None
    with path.open("rb") as stream:
        hasher = hashlib.sha256()
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            hasher.update(chunk)
        return hasher.hexdigest()


def write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(value, stream, indent=2, ensure_ascii=False, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(path)


def validate_workbook(path: Path) -> None:
    # Read only: saving with openpyxl here would erase Excel's formula caches again.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if not wb.sheetnames:
            raise ValueError("Workbook has no worksheets")
        for ws in wb:
            for _ in ws.iter_rows(values_only=True):
                pass
    finally:
        wb.close()


class ExcelClient:
    """Client for Spreadsheet-RL's /recalculate submit+poll protocol.

    Only predicted workbooks leave the machine. Job IDs survive interruptions;
    network/server errors leave cases pending, rather than turning them into zeros.
    """

    def __init__(self, base_url: str, *, engine_label: str, job_timeout: float = 600,
                 poll_interval: float = 1, max_bytes: int = 100 * 1024 * 1024,
                 transport=None):
        parts = urlsplit(base_url)
        if (parts.scheme not in {"http", "https"} or not parts.netloc or
                parts.username or parts.password or parts.query or parts.fragment):
            raise ValueError("Use an HTTP(S) base URL without credentials, query or fragment")
        if job_timeout <= 0 or poll_interval < 0 or max_bytes <= 0:
            raise ValueError("Invalid timeout, poll interval or file size limit")
        self.base_url = base_url.rstrip("/")
        self.engine_label = engine_label
        self.job_timeout = job_timeout
        self.poll_interval = poll_interval
        self.max_bytes = max_bytes
        token = os.environ.get("EXCEL_API_TOKEN")
        headers = {"Authorization": f"Bearer {token}"} if token else {}
        self.http = httpx.Client(headers=headers, timeout=40, follow_redirects=False,
                                 transport=transport)

    def close(self):
        self.http.close()

    def health(self) -> dict:
        response = self.http.get(self.base_url + "/health")
        if response.status_code != 200:
            raise RecalculationError(f"Excel service health HTTP {response.status_code}")
        value = response.json()
        if value.get("status") != "ok" or value.get("ready") is False:
            raise RecalculationError("Excel service is not ready")
        return value

    def recalculate(self, source: Path, target: Path, state_path: Path,
                    *, retry_failed: bool = False) -> Path:
        if source.resolve() == target.resolve():
            raise ValueError("Recalculated file must not overwrite the prediction")
        identity = {"source_sha256": digest(source), "base_url": self.base_url,
                    "engine_label": self.engine_label}
        if identity["source_sha256"] is None:
            raise RecalculationError(f"Missing prediction: {source}")
        state = json.loads(state_path.read_text()) if state_path.exists() else {}
        if state and state.get("identity") != identity:
            raise RecalculationError("Source/service changed; use a new report directory")
        if state.get("status") == "done":
            if digest(target) != state.get("output_sha256"):
                raise RecalculationError("Recalculated output missing or changed")
            return target
        if state.get("status") == "failed":
            if not retry_failed:
                raise RecalculationError("Remote job failed/expired; inspect state then use --retry-failed")
            state.setdefault("previous_jobs", []).append({
                "job_id": state.get("job_id"), "error": state.get("error")})
            state.pop("job_id", None)
        state.update(identity=identity, status="pending")
        state.pop("error", None)
        write_json(state_path, state)
        deadline = time.monotonic() + self.job_timeout
        try:
            while time.monotonic() < deadline:
                remaining = max(0.01, deadline - time.monotonic())
                try:
                    if not state.get("job_id"):
                        if source.stat().st_size > self.max_bytes:
                            raise RecalculationError("Prediction exceeds upload size limit")
                        with source.open("rb") as stream:
                            response = self.http.post(
                                self.base_url + "/recalculate/submit",
                                files={"file": (source.name, stream,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                                timeout=min(40, remaining))
                        if response.status_code in {429, 502, 503, 504}:
                            time.sleep(min(self.poll_interval, remaining))
                            continue
                        if response.status_code != 200:
                            raise RecalculationError(f"Submit HTTP {response.status_code}")
                        payload = response.json()
                        if payload.get("status") == "busy":
                            time.sleep(min(self.poll_interval, remaining))
                            continue
                        job_id = payload.get("job_id", "")
                        if not re.fullmatch(r"[A-Za-z0-9_-]+", job_id):
                            raise RecalculationError("Invalid job ID from Excel service")
                        state.update(job_id=job_id, status="submitted")
                        # Persist before polling. Never trust a service-supplied foreign URL.
                        write_json(state_path, state)
                    wait_s = min(10, max(0, deadline - time.monotonic()))
                    url = self.base_url + f"/recalculate/result/{state['job_id']}"
                    with self.http.stream("GET", url, params={"wait_s": wait_s},
                                          timeout=min(40, max(0.01, deadline - time.monotonic()))) as response:
                        if response.status_code in {429, 502, 503, 504}:
                            pass
                        elif response.status_code in {404, 410}:
                            state["status"] = "failed"
                            raise RecalculationError("Remote job expired or unknown")
                        elif response.status_code != 200:
                            raise RecalculationError(f"Result HTTP {response.status_code}")
                        elif "json" in response.headers.get("content-type", ""):
                            response.read()
                            payload = response.json()
                            if payload.get("status") == "error":
                                state["status"] = "failed"
                                raise RecalculationError(str(payload.get("msg", "Excel job failed"))[:2000])
                            if payload.get("status") not in {"queued", "running"}:
                                raise RecalculationError("Unexpected result status")
                        else:
                            target.parent.mkdir(parents=True, exist_ok=True)
                            temporary = target.with_name(target.stem + ".partial.xlsx")
                            try:
                                size = 0
                                with temporary.open("wb") as stream:
                                    for chunk in response.iter_bytes():
                                        size += len(chunk)
                                        if size > self.max_bytes:
                                            raise RecalculationError("Downloaded workbook exceeds size limit")
                                        stream.write(chunk)
                                    stream.flush()
                                    os.fsync(stream.fileno())
                                validate_workbook(temporary)
                                if digest(source) != identity["source_sha256"]:
                                    raise RecalculationError("Prediction changed during recalculation")
                                temporary.replace(target)
                            finally:
                                temporary.unlink(missing_ok=True)
                            state.update(status="done", output_sha256=digest(target))
                            write_json(state_path, state)
                            return target
                except (httpx.TimeoutException, httpx.NetworkError):
                    # A lost submit reply can orphan a remote job. No fabricated score is used.
                    state["last_transport_error"] = "network/timeout; submission may have been accepted"
                time.sleep(min(self.poll_interval, max(0, deadline - time.monotonic())))
            raise RecalculationError("Excel job timed out; rerun to poll the saved job ID")
        except Exception as exc:
            state["error"] = f"{type(exc).__name__}: {exc}"[:2000]
            write_json(state_path, state)
            raise RecalculationError(state["error"]) from exc
