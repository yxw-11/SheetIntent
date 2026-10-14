from __future__ import annotations
import ast
import json
import re
from collections import Counter
import openpyxl
from spreadsheet_code_agent_v0.prompts import extract_transform_code

def visible_facts(source, output, cell_budget=4096, max_changes=64):
    """Bounded deterministic sample. No task, answer path or scoring range enters here."""
    if cell_budget < 1 or max_changes < 1:
        raise ValueError("Observation budgets must be positive")
    before = openpyxl.load_workbook(source, data_only=False)
    try:
        after = openpyxl.load_workbook(output, data_only=False)
        try:
            result = {"input_sheets": before.sheetnames, "output_sheets": after.sheetnames,
                      "sheets": [], "sampled_cells": 0,
                      "scope": "Bounded sample, not a semantic verdict; formulas are not recalculated."}
            names = list(dict.fromkeys(before.sheetnames + after.sheetnames))[:min(8, cell_budget)]
            per_sheet = max(1, cell_budget // max(1, min(len(names), 8)))
            for name in names[:8]:
                a = before[name] if name in before else None
                b = after[name] if name in after else None
                rows = max(a.max_row if a else 0, b.max_row if b else 0)
                cols = min(32, max(a.max_column if a else 0, b.max_column if b else 0))
                before_shape = [a.max_row, a.max_column] if a else None
                after_shape = [b.max_row, b.max_column] if b else None
                selected_rows = list(range(1, min(rows, max(1, per_sheet // max(cols, 1) - 8)) + 1))
                selected_rows += [r for r in range(max(1, rows - 7), rows + 1) if r not in selected_rows]
                changes, headers, formulas, profiles, values = [], [], [], {}, []
                seen = 0
                for r in selected_rows:
                    for c in range(1, cols + 1):
                        if seen >= per_sheet:
                            break
                        seen += 1
                        av = a.cell(r, c).value if a else None
                        bv = b.cell(r, c).value if b else None
                        coord = openpyxl.utils.get_column_letter(c) + str(r)
                        profiles.setdefault(str(c), Counter())[type(av).__name__] += 1
                        result["sampled_cells"] += 1
                        small = lambda value: str(value)[:160] if value is not None else None
                        if av is not None or bv is not None:
                            values.append({"cell": coord,
                                "before": {"type": type(av).__name__, "value": small(av)},
                                "after": {"type": type(bv).__name__, "value": small(bv)}})
                        if r <= 3 and av is not None and len(headers) < 40:
                            headers.append([coord, small(av)])
                        if (type(av), av) != (type(bv), bv) and len(changes) < max_changes:
                            changes.append([coord, small(av), small(bv)])
                        if isinstance(bv, str) and bv.startswith("=") and len(formulas) < 12:
                            formulas.append([coord, small(bv)])
                # Retain unchanged input values too: headers and output differences
                # alone cannot ground arithmetic/type errors in the current program.
                shown = min(len(values), max_changes)
                samples = [values[i * (len(values) - 1) // max(1, shown - 1)] for i in range(shown)]
                result["sheets"].append({"name": name,
                    "before_shape": before_shape,
                    "after_shape": after_shape,
                    "headers": headers, "sampled_changes": changes, "output_formulas": formulas,
                    "sampled_values": samples, "omitted_scanned_values": len(values) - shown,
                    "input_types": {c: dict(v) for c, v in profiles.items()},
                    "merged_after": [str(x) for x in list(b.merged_cells.ranges)[:16]] if b else []})
            return result
        finally:
            after.close()
    finally:
        before.close()


def correction_prompt(instruction, code, facts):
    return (
        "Review the spreadsheet program against the user's instruction. Execution alone does not imply correctness.\n"
        "Use only the supplied input/output observations. Samples are incomplete; missing evidence is not proof of an error.\n"
        "Decide whether the existing program should be kept or revised. Preserve unrelated logic and cells.\n"
        "Return DECISION: KEEP with one short evidence sentence if no justified correction is needed.\n"
        "Otherwise return DECISION: REVISE, one short evidence sentence, and a complete Python code block defining transform(wb).\n"
        "The harness loads and saves the workbook. Do not read files, load/save workbooks, or use network access.\n"
        "openpyxl does not calculate formulas; write computed values when the instruction asks for resulting values.\n"
        f"\nINSTRUCTION\n{instruction}\n\nCURRENT PROGRAM\n```python\n{code}\n```\n"
        f"\nVISIBLE INPUT/OUTPUT OBSERVATIONS\n{json.dumps(facts, ensure_ascii=True, separators=(',', ':'))}"
    )


def decode(text, original=None):
    content = text.split("</think>")[-1].strip()
    if original is None:
        return "GENERATE", extract_transform_code(content)
    match = re.match(r"DECISION:\s*(KEEP|REVISE)\b", content)
    if not match:
        raise ValueError("Missing DECISION: KEEP or DECISION: REVISE")
    action = match.group(1)
    if action == "KEEP":
        if "```" in content:
            proposed = extract_transform_code(content)
            if ast.dump(ast.parse(proposed)) != ast.dump(ast.parse(original)):
                raise ValueError("KEEP must not change the program")
        return action, original
    code = extract_transform_code(content)
    if not re.search(r"(?m)^def\s+transform\s*\(", code):
        raise ValueError("REVISE must supply complete transform(wb) code")
    return action, code
