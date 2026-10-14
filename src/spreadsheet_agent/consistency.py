from __future__ import annotations
import hashlib
import json
import zipfile
from collections import Counter
from datetime import date, datetime, time, timedelta
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from .io import digest

def cell_value(value):
    if isinstance(value, ArrayFormula):
        return ["array_formula", value.ref, value.text]
    if isinstance(value, DataTableFormula):
        return ["data_table_formula", sorted(dict(value).items())]
    if isinstance(value, (datetime, date, time)):
        return [type(value).__name__, value.isoformat()]
    if isinstance(value, timedelta):
        return ["timedelta", value.total_seconds()]
    if value is None or isinstance(value, (str, bool, int, float)):
        return [type(value).__name__, value]
    raise ValueError(f"Unsupported cell value: {type(value).__name__}")


def fingerprint(path, max_cells=200000, max_xml_bytes=67108864):
    """Stream values/formulas, not XLSX bytes or absent formula caches. No recalculation.

    Bounds fail closed to 'not fingerprinted'; partial workbooks never vote.
    Styles and trailing empty cells are intentionally excluded from this proxy.
    """
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > max_xml_bytes:
            raise ValueError("Workbook exceeds fingerprint XML budget")
    wb = load_workbook(path, read_only=True, data_only=False)
    h = hashlib.sha256()
    def add(value):
        h.update((json.dumps(value, ensure_ascii=True, allow_nan=False) + "\n").encode())
    count = 0
    try:
        add(wb.sheetnames)
        for ws in wb:
            # Also count cells during iteration: malformed dimensions cannot bypass the cap.
            if ws.max_row and ws.max_column and count + ws.max_row * ws.max_column > max_cells:
                raise ValueError("Workbook exceeds fingerprint cell budget")
            add(["sheet", ws.title])
            for row in ws.iter_rows():
                for cell in row:
                    count += 1
                    if count > max_cells:
                        raise ValueError("Workbook exceeds fingerprint cell budget")
                    if cell.value is not None:
                        add([cell.coordinate, cell.data_type, cell_value(cell.value)])
        return h.hexdigest()
    finally:
        wb.close()


def select(observations):
    if not observations:
        raise ValueError("Empty candidate pool")
    valid = [i for i, o in enumerate(observations) if o["executable"]]
    first = valid[0] if valid else 0
    voters = [i for i in valid if observations[i]["fingerprint"] is not None]
    counts = Counter(observations[i]["fingerprint"] for i in voters)
    # A singleton is not consensus. Ties between nontrivial groups use generation order.
    if not counts or max(counts.values()) < 2:
        consensus = first
    else:
        consensus = min(i for i in voters if counts[observations[i]["fingerprint"]] == max(counts.values()))
    return {"first": 0, "first_executable": first, "consensus": consensus}

def strict_fingerprint(path):
    """Conservative agreement proxy: values/formulas plus stored formatting/merges.

    Style-table reordering can produce false negatives, never claimed semantic
    equivalence. Missing formula caches do not become matching empty answers.
    """
    value = fingerprint(path)
    extra = []
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        if "xl/styles.xml" in archive.namelist():
            extra.append(archive.read("xl/styles.xml").decode("utf-8"))
        for name in sorted(archive.namelist()):
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                sheet = ET.fromstring(archive.read(name))
                styles = [(c.get("r"), c.get("s")) for c in sheet.iter(ns + "c") if c.get("s")]
                layout = [ET.tostring(e, encoding="unicode") for e in sheet if e.tag in
                          {ns + s for s in ("mergeCells", "cols", "conditionalFormatting", "autoFilter")}]
                extra.append([name, styles, layout])
    return digest([value, extra])
