from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from spreadsheetbench_lite.evaluate_v1 import expand_cell_range, parse_answer_positions


def inspect_workbook(path: Path, understanding: dict[str, Any] | None, answer_position: str, rows: int = 8, max_cols: int = 20) -> dict[str, Any]:
    """Build compact workbook grounding without dumping the whole file."""

    wb = load_workbook(path, data_only=False)
    text_hint = " ".join(str((understanding or {}).get(k, "")) for k in ("where", "from", "what"))
    selected_sheets = _select_sheets(wb, text_hint, answer_position)
    info: dict[str, Any] = {
        "sheet_names": wb.sheetnames,
        "selected_sheets": selected_sheets,
        "target_ranges": _safe_answer_ranges(wb, answer_position),
        "target_context": _target_context(wb, answer_position),
        "sheets": [],
    }
    for ws in wb.worksheets:
        sheet_summary: dict[str, Any] = {
            "name": ws.title,
            "dimensions": {"rows": ws.max_row, "cols": ws.max_column},
            "merged_ranges": [str(rng) for rng in list(ws.merged_cells.ranges)[:20]],
        }
        if ws.title in selected_sheets:
            sheet_summary["preview_rows"] = _preview_rows(ws, rows, max_cols)
            sheet_summary["headers"] = _header_candidates(ws, rows, max_cols)
            sheet_summary["formulas"] = _formulas(ws, limit=25)
            sheet_summary["type_profile"] = _type_profile(ws, rows=min(ws.max_row, 120), max_cols=max_cols)
            sheet_summary["range_samples"] = _range_samples(ws, text_hint + " " + answer_position)
        info["sheets"].append(sheet_summary)
    return info


def render_inspection(info: dict[str, Any], limit: int = 9000) -> str:
    return json.dumps(info, ensure_ascii=False, indent=2, default=str)[:limit]


def _select_sheets(wb, hint: str, answer_position: str) -> list[str]:
    selected = []
    lowered = hint.lower()
    for name in wb.sheetnames:
        if name.lower() in lowered:
            selected.append(name)
    try:
        for sheet_name, _ in parse_answer_positions(wb, answer_position):
            if sheet_name in wb.sheetnames and sheet_name not in selected:
                selected.append(sheet_name)
    except Exception:
        pass
    if wb.sheetnames and wb.sheetnames[0] not in selected:
        selected.insert(0, wb.sheetnames[0])
    return selected[:4]


def _safe_answer_ranges(wb, answer_position: str) -> list[dict[str, str]]:
    out = []
    try:
        for sheet, rng in parse_answer_positions(wb, answer_position):
            out.append({"sheet": sheet, "range": rng})
    except Exception as exc:
        out.append({"error": str(exc), "raw": answer_position})
    return out


def _target_context(wb, answer_position: str) -> list[dict[str, Any]]:
    contexts = []
    try:
        positions = parse_answer_positions(wb, answer_position)
    except Exception:
        return contexts
    for sheet_name, cell_range in positions[:6]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        try:
            cells = expand_cell_range(cell_range, ws)
        except Exception:
            continue
        if not cells:
            continue
        first = cells[0]
        match = re.match(r"([A-Z]+)(\d+)$", first)
        if not match:
            continue
        col_letters, row_text = match.groups()
        row = int(row_text)
        col = 0
        for ch in col_letters:
            col = col * 26 + ord(ch) - ord("A") + 1
        contexts.append(
            {
                "sheet": sheet_name,
                "range": cell_range,
                "target_sample": _cell_values(ws, cells[:20]),
                "nearby_values": _nearby_values(ws, row, col),
                "same_row_nonempty": _same_row_nonempty(ws, row),
                "same_col_nonempty": _same_col_nonempty(ws, col),
            }
        )
    return contexts


def _cell_values(ws, cells: list[str]) -> list[dict[str, str]]:
    values = []
    for cell_name in cells:
        cell = ws[cell_name]
        if cell.value not in (None, ""):
            values.append({"cell": cell_name, "value": _fmt(cell.value), "type": _kind(cell.value)})
    return values[:20]


def _nearby_values(ws, row: int, col: int) -> list[list[str]]:
    rows = []
    min_row = max(1, row - 2)
    max_row = min(ws.max_row, row + 2)
    min_col = max(1, col - 2)
    max_col = min(ws.max_column, col + 2)
    for r in range(min_row, max_row + 1):
        rows.append([_fmt(ws.cell(r, c).value, 80) for c in range(min_col, max_col + 1)])
    return rows


def _same_row_nonempty(ws, row: int, limit: int = 12) -> list[dict[str, str]]:
    out = []
    if row < 1 or row > ws.max_row:
        return out
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row, col).value
        if value not in (None, ""):
            out.append({"cell": f"{get_column_letter(col)}{row}", "value": _fmt(value, 80)})
            if len(out) >= limit:
                break
    return out


def _same_col_nonempty(ws, col: int, limit: int = 12) -> list[dict[str, str]]:
    out = []
    if col < 1 or col > ws.max_column:
        return out
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, col).value
        if value not in (None, ""):
            out.append({"cell": f"{get_column_letter(col)}{row}", "value": _fmt(value, 80)})
            if len(out) >= limit:
                break
    return out


def _preview_rows(ws, rows: int, max_cols: int) -> list[list[str]]:
    rendered = []
    for row in ws.iter_rows(max_row=min(rows, ws.max_row), max_col=min(max_cols, ws.max_column)):
        rendered.append([_fmt(cell.value) for cell in row])
    return rendered


def _header_candidates(ws, rows: int, max_cols: int) -> list[dict[str, Any]]:
    headers = []
    for row_idx in range(1, min(rows, ws.max_row) + 1):
        values = [_fmt(ws.cell(row_idx, col).value) for col in range(1, min(max_cols, ws.max_column) + 1)]
        if sum(bool(v) for v in values) >= 2:
            headers.append({"row": row_idx, "values": values})
    return headers[:6]


def _formulas(ws, limit: int) -> list[dict[str, str]]:
    formulas = []
    for row in ws.iter_rows(max_row=min(ws.max_row, 120), max_col=min(ws.max_column, 40)):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formulas.append({"cell": cell.coordinate, "formula": _fmt(value, 180)})
                if len(formulas) >= limit:
                    return formulas
    return formulas


def _type_profile(ws, rows: int, max_cols: int) -> list[dict[str, Any]]:
    profile = []
    for col in range(1, min(max_cols, ws.max_column) + 1):
        counts: Counter[str] = Counter()
        examples: dict[str, str] = {}
        for row in range(1, rows + 1):
            value = ws.cell(row, col).value
            kind = _kind(value)
            counts[kind] += 1
            if kind != "blank" and kind not in examples:
                examples[kind] = _fmt(value, 80)
        profile.append({"col": get_column_letter(col), "types": dict(counts), "examples": examples})
    return profile


def _range_samples(ws, text: str) -> list[dict[str, Any]]:
    samples = []
    for token in _range_tokens(text):
        try:
            cells = expand_cell_range(token)
        except Exception:
            continue
        values = []
        for cell_name in cells[:40]:
            try:
                values.append({"cell": cell_name, "value": _fmt(ws[cell_name].value)})
            except Exception:
                pass
        if values:
            samples.append({"range": token, "sample": values})
    return samples[:8]


def _range_tokens(text: str) -> list[str]:
    tokens = []
    for match in re.finditer(r"\b[A-Za-z]{1,3}\d+(?::[A-Za-z]{1,3}\d+)?\b", text):
        token = match.group(0).upper()
        if token not in tokens:
            tokens.append(token)
    return tokens[:12]


def _kind(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number"
    if hasattr(value, "year") and hasattr(value, "month"):
        return "date"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, str):
        try:
            float(value.replace(",", ""))
            return "numeric_string"
        except ValueError:
            return "string"
    return type(value).__name__


def _fmt(value: Any, limit: int = 120) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]
