from __future__ import annotations

import inspect
from copy import copy
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import cell as cell_utils
from openpyxl.worksheet.worksheet import Worksheet


BASE_APIS = [
    ("Workbook.create_sheet", Workbook.create_sheet, "Create a worksheet by title."),
    ("Workbook.copy_worksheet", Workbook.copy_worksheet, "Copy a worksheet inside the same workbook."),
    ("Worksheet.cell", Worksheet.cell, "Access a cell by 1-based row and column integers."),
    ("Worksheet.iter_rows", Worksheet.iter_rows, "Iterate rows; returns tuples of Cell objects unless values_only=True."),
    ("Worksheet.max_row", None, "Property: last used row index."),
    ("Worksheet.max_column", None, "Property: last used column index; there is no max_col."),
    ("Cell.value", None, "Cell value property."),
]

KEYWORD_APIS = {
    "delete": [("Worksheet.delete_rows", Worksheet.delete_rows, "delete_rows(idx, amount=1) deletes whole rows."), ("Worksheet.delete_cols", Worksheet.delete_cols, "delete_cols(idx, amount=1) deletes whole columns.")],
    "insert": [("Worksheet.insert_rows", Worksheet.insert_rows, "insert_rows(idx, amount=1)."), ("Worksheet.insert_cols", Worksheet.insert_cols, "insert_cols(idx, amount=1).")],
    "iterate": [("Worksheet.iter_rows", Worksheet.iter_rows, "Use for row iteration. With values_only=True it returns values, not Cell objects.")],
    "row": [("Worksheet.iter_rows", Worksheet.iter_rows, "Rows are tuples; iterate cells inside the tuple.")],
    "column": [("openpyxl.utils.get_column_letter", cell_utils.get_column_letter, "Convert 1-based column index to letters."), ("openpyxl.utils.column_index_from_string", cell_utils.column_index_from_string, "Convert column letters to 1-based index.")],
    "range": [("openpyxl.utils.range_boundaries", cell_utils.range_boundaries, "Parse A1 range into min_col, min_row, max_col, max_row."), ("Worksheet.move_range", Worksheet.move_range, "Move a cell range by rows/cols offsets.")],
    "merge": [("Worksheet.merge_cells", Worksheet.merge_cells, "Merge cells; only top-left merged cell is writable."), ("Worksheet.unmerge_cells", Worksheet.unmerge_cells, "Unmerge a merged range before writing inner cells.")],
    "format": [("Cell.font", None, "Use cell.font = copy(other.font); do not use cell.bold."), ("Cell.fill", None, "Use cell.fill = copy(other.fill)."), ("Cell.number_format", None, "Assign number format strings directly.")],
    "style": [("copy.copy", copy, "Use copy(style_object) before assigning styles between cells."), ("openpyxl.styles.Font", Font, "Create font styles, e.g. Font(bold=True).")],
    "formula": [("Cell.value", None, "Formulas are strings beginning with '='. openpyxl does not calculate formulas.")],
    "append": [("Worksheet.append", Worksheet.append, "Append values/lists, not Cell objects from another sheet.")],
    "copy": [("copy.copy", copy, "Use copy for style objects."), ("Workbook.copy_worksheet", Workbook.copy_worksheet, "Copy a worksheet inside the same workbook.")],
}

CANONICAL_EXAMPLES = [
    "for row in ws.iter_rows(min_row=2, max_row=ws.max_row):\n    values = [cell.value for cell in row]",
    "ws.delete_rows(idx=row_index, amount=1)  # no columns= argument",
    "ws.cell(row=r, column=c).value = value",
    "target.font = copy(source.font); target.fill = copy(source.fill)",
]


def lookup_openpyxl_apis(api_needs: list[str] | None, limit: int = 14) -> list[dict[str, str]]:
    selected = list(BASE_APIS)
    needs = " ".join(api_needs or []).lower()
    for key, apis in KEYWORD_APIS.items():
        if key in needs:
            selected.extend(apis)
    if not needs:
        selected.extend(KEYWORD_APIS["range"] + KEYWORD_APIS["format"] + KEYWORD_APIS["delete"])
    seen = set()
    refs = []
    for name, obj, note in selected:
        if name in seen:
            continue
        seen.add(name)
        refs.append({"api": name, "signature": _signature(obj), "note": note, "doc": _doc(obj)})
        if len(refs) >= limit:
            break
    return refs


def render_api_refs(refs: list[dict[str, str]]) -> str:
    lines = []
    for ref in refs:
        lines.append(f"- {ref['api']}{ref['signature']}: {ref['note']}")
        if ref.get("doc"):
            lines.append(f"  doc: {ref['doc']}")
    if refs:
        lines.append("Canonical examples:")
        lines.extend(f"```python\n{example}\n```" for example in CANONICAL_EXAMPLES)
    return "\n".join(lines)


def _signature(obj: Any) -> str:
    if obj is None:
        return ""
    try:
        return str(inspect.signature(obj))
    except Exception:
        return ""


def _doc(obj: Any, limit: int = 180) -> str:
    if obj is None:
        return ""
    doc = inspect.getdoc(obj) or ""
    doc = " ".join(doc.split())
    return doc[:limit]
