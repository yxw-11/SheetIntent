from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from spreadsheetbench_lite.evaluate_v1 import expand_cell_range, parse_answer_positions


@dataclass
class AgentExecResult:
    ok: bool
    returncode: int
    stdout: str
    stderr: str
    timed_out: bool = False
    error_category: str = ""
    verification: dict[str, Any] = field(default_factory=dict)


def run_transform_code(code: str, input_path: Path, output_path: Path, work_dir: Path, timeout: int, answer_position: str) -> AgentExecResult:
    work_dir.mkdir(parents=True, exist_ok=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    solution_file = work_dir / "solution.py"
    runner_file = work_dir / "run_transform.py"
    solution_file.write_text(code, encoding="utf-8")
    runner_file.write_text(_runner_source(input_path, output_path), encoding="utf-8")
    env = os.environ.copy()
    env.setdefault("PYTHONNOUSERSITE", "1")
    try:
        proc = subprocess.run(
            [sys.executable, str(runner_file)],
            cwd=work_dir,
            env=env,
            text=True,
            capture_output=True,
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        return AgentExecResult(False, -1, exc.stdout or "", exc.stderr or f"Timed out after {timeout} seconds.", True, "timeout")

    verification = verify_output(input_path, output_path, answer_position) if proc.returncode == 0 else {}
    ok = proc.returncode == 0 and verification.get("openable", False)
    return AgentExecResult(ok, proc.returncode, proc.stdout, proc.stderr, False, classify_runtime_error(proc.stderr), verification)


def verify_output(input_path: Path, output_path: Path, answer_position: str) -> dict[str, Any]:
    result: dict[str, Any] = {
        "exists": output_path.exists(),
        "openable": False,
        "changed_target_values": None,
        "target_all_empty": None,
        "target_total": 0,
        "target_empty": 0,
        "target_formula_cells": 0,
        "formula_cache_risk_cells": [],
        "excel_error_values": [],
    }
    if not output_path.exists():
        return result
    try:
        input_wb = openpyxl.load_workbook(input_path, data_only=False)
        output_wb = openpyxl.load_workbook(output_path, data_only=False)
        output_values_wb = openpyxl.load_workbook(output_path, data_only=True)
    except Exception as exc:
        result["open_error"] = str(exc)
        return result
    result["openable"] = True
    try:
        if answer_position in ("", "__FULL_WORKBOOK__"):
            positions = [(name, output_wb[name].dimensions) for name in output_wb.sheetnames]
        else:
            positions = parse_answer_positions(output_wb, answer_position)
    except Exception as exc:
        result["position_error"] = str(exc)
        return result

    total = 0
    empty = 0
    changed = 0
    formula_cells = 0
    formula_cache_risks = []
    errors = []
    for sheet_name, cell_range in positions:
        if sheet_name not in output_wb.sheetnames or sheet_name not in input_wb.sheetnames:
            continue
        in_ws = input_wb[sheet_name]
        out_ws = output_wb[sheet_name]
        out_values_ws = output_values_wb[sheet_name]
        try:
            cells = expand_cell_range(cell_range, out_ws)
        except Exception:
            continue
        for cell_name in cells[:5000]:
            total += 1
            out_value = out_ws[cell_name].value
            eval_value = out_values_ws[cell_name].value
            in_value = in_ws[cell_name].value
            if eval_value in (None, ""):
                empty += 1
            if out_value != in_value:
                changed += 1
            if isinstance(out_value, str) and out_value.startswith("="):
                formula_cells += 1
                if eval_value in (None, ""):
                    formula_cache_risks.append(f"{sheet_name}!{cell_name}")
            if isinstance(out_value, str) and out_value.startswith("#"):
                errors.append({"cell": f"{sheet_name}!{cell_name}", "value": out_value})
                if len(errors) >= 20:
                    break
    result["changed_target_values"] = changed
    result["target_total"] = total
    result["target_empty"] = empty
    result["target_formula_cells"] = formula_cells
    result["target_all_empty"] = bool(total and empty == total)
    result["formula_cache_risk_cells"] = formula_cache_risks[:20]
    result["excel_error_values"] = errors
    return result


def classify_runtime_error(stderr: str) -> str:
    text = stderr or ""
    if not text:
        return ""
    if "SyntaxError" in text or "IndentationError" in text:
        return "syntax_error"
    if "AttributeError" in text:
        if "MergedCell" in text and "read-only" in text:
            return "merged_cell_error"
        return "attribute_error"
    if "TypeError" in text:
        return "type_error"
    if "ValueError" in text:
        if "coordinate" in text or "range" in text or "Cell coordinates" in text:
            return "range_error"
        return "value_error"
    if "IndexError" in text:
        return "index_error"
    if "KeyError" in text or ("Worksheet" in text and "does not exist" in text):
        return "sheet_error"
    if "FileNotFoundError" in text:
        return "file_error"
    if "Timed out" in text or "Timeout" in text:
        return "timeout"
    return "other_runtime_error"


def runtime_diagnostic(result: AgentExecResult, code: str) -> dict[str, Any]:
    return {
        "error_type": result.error_category or "runtime_error",
        "message": _last_error_line(result.stderr),
        "traceback_tail": result.stderr[-3000:],
        "offending_code": _near_line(code, result.stderr),
        "verification": result.verification,
    }


def preflight_diagnostic(preflight: dict[str, Any]) -> dict[str, Any]:
    first = (preflight.get("errors") or [{}])[0]
    return {
        "error_type": first.get("category", "preflight_error"),
        "message": first.get("message", "preflight failed"),
        "offending_code": first.get("snippet", ""),
        "preflight": preflight,
    }


def _runner_source(input_path: Path, output_path: Path) -> str:
    return f"""from pathlib import Path
from openpyxl import load_workbook

input_path = {json.dumps(str(input_path))}
output_path = {json.dumps(str(output_path))}
Path(output_path).parent.mkdir(parents=True, exist_ok=True)

namespace = {{"__file__": str(Path("solution.py").resolve()), "__name__": "__generated_transform__"}}
exec(Path("solution.py").read_text(encoding="utf-8"), namespace)
transform = namespace.get("transform")
if not callable(transform):
    raise RuntimeError("solution must define transform(wb) -> None")
wb = load_workbook(input_path)
transform(wb)
wb.save(output_path)
if not Path(output_path).exists():
    raise RuntimeError(f"solution did not create output file: {{output_path}}")
"""


def _last_error_line(stderr: str) -> str:
    for line in reversed((stderr or "").splitlines()):
        line = line.strip()
        if line:
            return line[:500]
    return ""


def _near_line(code: str, stderr: str) -> str:
    match = re.search(r"line (\d+)", stderr or "")
    if not match:
        return "\n".join(code.splitlines()[:20])[:1000]
    line_no = int(match.group(1))
    lines = code.splitlines()
    start = max(0, line_no - 4)
    end = min(len(lines), line_no + 3)
    return "\n".join(f"{idx + 1}: {lines[idx]}" for idx in range(start, end))[:1000]
